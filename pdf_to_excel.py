#!/usr/bin/env python3
"""Extract tables and text from a PDF into an Excel workbook.

Usage:
    python pdf_to_excel.py input.pdf output.xlsx
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, List, Optional

import pdfplumber
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


MAX_SHEET_NAME_LEN = 31


def sanitize_sheet_name(name: str) -> str:
    """Return a valid Excel sheet name."""
    invalid_chars = set('[]:*?/\\')
    cleaned = "".join("_" if ch in invalid_chars else ch for ch in name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:MAX_SHEET_NAME_LEN]


def autosize_columns(worksheet) -> None:
    """Set worksheet column widths based on longest cell content."""
    max_lengths = {}
    for row in worksheet.iter_rows(values_only=True):
        for idx, value in enumerate(row, start=1):
            if value is None:
                continue
            length = len(str(value))
            max_lengths[idx] = max(max_lengths.get(idx, 0), length)

    for idx, length in max_lengths.items():
        worksheet.column_dimensions[get_column_letter(idx)].width = min(max(length + 2, 10), 60)


def normalize_row(row: Iterable[Optional[str]]) -> List[str]:
    """Normalize table row cells by replacing None and trimming whitespace."""
    normalized: List[str] = []
    for cell in row:
        if cell is None:
            normalized.append("")
        else:
            normalized.append(" ".join(str(cell).split()))
    return normalized


def write_table_sheet(workbook: Workbook, page_number: int, table_index: int, table_rows: List[List[str]]) -> None:
    sheet_name = sanitize_sheet_name(f"P{page_number}_Table{table_index}")
    ws = workbook.create_sheet(title=sheet_name)

    max_cols = max((len(r) for r in table_rows), default=0)
    for row in table_rows:
        padded = row + [""] * (max_cols - len(row))
        ws.append(padded)

    autosize_columns(ws)


def write_text_sheet(workbook: Workbook, page_text_rows: List[List[str]]) -> None:
    ws = workbook.create_sheet(title="Extracted_Text")
    ws.append(["Page", "Line"])
    for row in page_text_rows:
        ws.append(row)
    autosize_columns(ws)


def extract_pdf_to_excel(pdf_path: Path, excel_path: Path) -> dict:
    """Extract PDF tables and text into an Excel workbook.

    Returns a summary dictionary with extraction stats.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    table_count = 0
    text_lines = 0
    page_text_rows: List[List[str]] = []

    page_total = 0

    with pdfplumber.open(str(pdf_path)) as pdf:
        page_total = len(pdf.pages)
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []

            for table_idx, table in enumerate(tables, start=1):
                cleaned_rows = [normalize_row(row) for row in table if row]
                if cleaned_rows:
                    write_table_sheet(workbook, page_num, table_idx, cleaned_rows)
                    table_count += 1

            text = page.extract_text() or ""
            if text.strip():
                for line in text.splitlines():
                    stripped = line.strip()
                    if stripped:
                        page_text_rows.append([str(page_num), stripped])
                        text_lines += 1

    if page_text_rows:
        write_text_sheet(workbook, page_text_rows)

    if not workbook.sheetnames:
        ws = workbook.create_sheet(title="No_Data")
        ws.append(["No tables or text were extracted from the PDF."])
        autosize_columns(ws)

    workbook.save(str(excel_path))

    return {
        "pages": page_total,
        "tables": table_count,
        "text_lines": text_lines,
        "output": str(excel_path),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract tables/text from PDF to Excel.")
    parser.add_argument("input_pdf", type=Path, help="Path to input PDF")
    parser.add_argument("output_xlsx", type=Path, help="Path to output XLSX")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if not args.input_pdf.exists():
        raise SystemExit(f"Input PDF does not exist: {args.input_pdf}")

    if args.input_pdf.suffix.lower() != ".pdf":
        raise SystemExit("Input file must be a PDF.")

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    summary = extract_pdf_to_excel(args.input_pdf, args.output_xlsx)
    print(
        "Extraction complete | "
        f"Pages: {summary['pages']} | Tables: {summary['tables']} | "
        f"Text lines: {summary['text_lines']} | Output: {summary['output']}"
    )


if __name__ == "__main__":
    main()
